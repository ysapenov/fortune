"""Tests for resume parser and keyword matcher."""

import io
from docx import Document
from app.services.resume_parser import parse_resume, _parse_text_to_structured
from app.services.keyword_matcher import compute_match, compute_semantic_match, rank_keywords


def test_resume_text_parser():
    sample_text = """
    Jane Smith
    jane.smith@example.com | (555) 987-6543 | New York, NY
    
    Summary
    Senior Software Engineer with 7+ years of experience in distributed systems.
    
    Experience
    Acme Corp | Lead Developer | 2020 - Present
    • Architected high-throughput microservices using Python and Go.
    • Managed a cluster of 50 Kubernetes nodes on AWS.
    
    Education
    New York University | Bachelor of Science in Computer Science | 2013 - 2017
    
    Skills
    Python, Go, AWS, Kubernetes, Docker, PostgreSQL, Redis
    
    Certifications
    AWS Certified Solutions Architect
    """

    structured = _parse_text_to_structured(sample_text)
    assert structured["email"] == "jane.smith@example.com"
    assert "555" in structured["phone"]
    assert "Jane Smith" in structured["name"]
    assert len(structured["experience"]) >= 1
    assert len(structured["skills"]) >= 3


def test_docx_parser():
    doc = Document()
    doc.add_heading("John Developer", level=1)
    doc.add_paragraph("john@example.com | 123-456-7890")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, FastAPI, Docker")
    
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()
    
    parsed = parse_resume(docx_bytes, "sample.docx")
    assert parsed["email"] == "john@example.com"
    assert len(parsed["skills"]) >= 1


def test_compute_match_and_semantic_match():
    resume_text = "Experienced in Python, Docker, Kubernetes, AWS."
    job_desc = "Seeking engineer with Python, Kubernetes, AWS, and GCP."
    
    res = compute_match(resume_text, job_desc)
    assert res["match_percentage"] == 75.0  # 3 of 4 keywords
    assert "gcp" in res["missing_keywords"]
    assert "python" in res["matched_keywords"]
    
    ranked = rank_keywords("Python Python Python Docker Docker AWS")
    assert ranked[0] == ("python", 3)
    assert ranked[1] == ("docker", 2)
    
    # Semantic match test with fallback or mocked embedding
    sem_res = compute_semantic_match(resume_text, job_desc)
    assert "keyword_match" in sem_res
    assert "combined_score" in sem_res
