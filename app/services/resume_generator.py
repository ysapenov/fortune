"""
Resume Generator Service

Generates PDF and DOCX files from structured resume JSON data.
Uses reportlab for PDF generation and python-docx for DOCX generation.
"""

import io
from typing import Dict, Any, List


def generate_pdf(resume_data: Dict[str, Any]) -> bytes:
    """
    Generate a professional PDF resume from structured data.
    
    Args:
        resume_data: Structured resume JSON
        
    Returns:
        PDF file content as bytes
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable, ListFlowable, ListItem
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    # Define styles
    styles = getSampleStyleSheet()

    name_style = ParagraphStyle(
        "NameStyle",
        parent=styles["Title"],
        fontSize=18,
        spaceAfter=4,
        textColor=HexColor("#1a1a2e"),
        leading=22,
    )

    contact_style = ParagraphStyle(
        "ContactStyle",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=8,
        textColor=HexColor("#555555"),
        alignment=1,  # Center
    )

    section_header_style = ParagraphStyle(
        "SectionHeader",
        parent=styles["Heading2"],
        fontSize=12,
        spaceBefore=12,
        spaceAfter=4,
        textColor=HexColor("#1a1a2e"),
        borderWidth=0,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=4,
        leading=13,
    )

    bold_style = ParagraphStyle(
        "BoldStyle",
        parent=body_style,
        fontName="Helvetica-Bold",
    )

    bullet_style = ParagraphStyle(
        "BulletStyle",
        parent=body_style,
        leftIndent=20,
        bulletIndent=10,
        spaceAfter=2,
    )

    story = []

    # Name
    name = resume_data.get("name", "")
    if name:
        story.append(Paragraph(name, name_style))

    # Contact info line
    contact_parts = []
    email = resume_data.get("email", "")
    phone = resume_data.get("phone", "")
    location = resume_data.get("location", "")
    if email:
        contact_parts.append(email)
    if phone:
        contact_parts.append(phone)
    if location:
        contact_parts.append(location)
    if contact_parts:
        story.append(Paragraph(" | ".join(contact_parts), contact_style))

    # Divider
    story.append(HRFlowable(width="100%", thickness=1, color=HexColor("#1a1a2e")))
    story.append(Spacer(1, 6))

    # Summary
    summary = resume_data.get("summary", "")
    if summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", section_header_style))
        story.append(Paragraph(summary, body_style))
        story.append(Spacer(1, 4))

    # Experience
    experience = resume_data.get("experience", [])
    if experience:
        story.append(Paragraph("EXPERIENCE", section_header_style))
        for exp in experience:
            title = exp.get("title", "")
            company = exp.get("company", "")
            dates = exp.get("dates", "")
            description = exp.get("description", "")

            # Title and company line
            header_text = ""
            if title and company:
                header_text = f"<b>{title}</b> — {company}"
            elif title:
                header_text = f"<b>{title}</b>"
            elif company:
                header_text = f"<b>{company}</b>"

            if dates:
                header_text += f"  <i>({dates})</i>"

            if header_text:
                story.append(Paragraph(header_text, body_style))

            # Description bullets
            if description:
                desc_lines = [l.strip() for l in description.split("\n") if l.strip()]
                for desc_line in desc_lines:
                    story.append(Paragraph(f"• {_escape_xml(desc_line)}", bullet_style))

            story.append(Spacer(1, 4))

    # Education
    education = resume_data.get("education", [])
    if education:
        story.append(Paragraph("EDUCATION", section_header_style))
        for edu in education:
            institution = edu.get("institution", "")
            degree = edu.get("degree", "")
            dates = edu.get("dates", "")
            gpa = edu.get("gpa", "")

            edu_text = ""
            if degree and institution:
                edu_text = f"<b>{_escape_xml(degree)}</b> — {_escape_xml(institution)}"
            elif institution:
                edu_text = f"<b>{_escape_xml(institution)}</b>"
            elif degree:
                edu_text = f"<b>{_escape_xml(degree)}</b>"

            if dates:
                edu_text += f"  <i>({_escape_xml(dates)})</i>"

            if edu_text:
                story.append(Paragraph(edu_text, body_style))

            if gpa:
                story.append(Paragraph(f"GPA: {_escape_xml(gpa)}", bullet_style))

        story.append(Spacer(1, 4))

    # Skills
    skills = resume_data.get("skills", [])
    if skills:
        story.append(Paragraph("SKILLS", section_header_style))
        skills_text = ", ".join(_escape_xml(s) for s in skills)
        story.append(Paragraph(skills_text, body_style))
        story.append(Spacer(1, 4))

    # Certifications
    certifications = resume_data.get("certifications", [])
    if certifications:
        story.append(Paragraph("CERTIFICATIONS", section_header_style))
        for cert in certifications:
            story.append(Paragraph(f"• {_escape_xml(cert)}", bullet_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buffer.getvalue()


def generate_docx(resume_data: Dict[str, Any]) -> bytes:
    """
    Generate an ATS-friendly DOCX resume from structured data.
    
    Uses simple formatting without tables or columns for maximum
    ATS compatibility.
    
    Args:
        resume_data: Structured resume JSON
        
    Returns:
        DOCX file content as bytes
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Configure default style
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Configure margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Name
    name = resume_data.get("name", "")
    if name:
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_para.add_run(name)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    # Contact info
    contact_parts = []
    email = resume_data.get("email", "")
    phone = resume_data.get("phone", "")
    location = resume_data.get("location", "")
    if email:
        contact_parts.append(email)
    if phone:
        contact_parts.append(phone)
    if location:
        contact_parts.append(location)
    if contact_parts:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = contact_para.add_run(" | ".join(contact_parts))
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Summary
    summary = resume_data.get("summary", "")
    if summary:
        _add_section_header(doc, "PROFESSIONAL SUMMARY")
        doc.add_paragraph(summary)

    # Experience
    experience = resume_data.get("experience", [])
    if experience:
        _add_section_header(doc, "EXPERIENCE")
        for exp in experience:
            title = exp.get("title", "")
            company = exp.get("company", "")
            dates = exp.get("dates", "")
            description = exp.get("description", "")

            # Title/company line
            exp_para = doc.add_paragraph()
            if title:
                run = exp_para.add_run(title)
                run.bold = True
            if company:
                if title:
                    exp_para.add_run(" — ")
                run = exp_para.add_run(company)
            if dates:
                exp_para.add_run(f"  ({dates})")

            # Description bullets
            if description:
                for line in description.split("\n"):
                    line = line.strip()
                    if line:
                        bullet_para = doc.add_paragraph(line, style="List Bullet")

    # Education
    education = resume_data.get("education", [])
    if education:
        _add_section_header(doc, "EDUCATION")
        for edu in education:
            institution = edu.get("institution", "")
            degree = edu.get("degree", "")
            dates = edu.get("dates", "")
            gpa = edu.get("gpa", "")

            edu_para = doc.add_paragraph()
            if degree:
                run = edu_para.add_run(degree)
                run.bold = True
            if institution:
                if degree:
                    edu_para.add_run(" — ")
                edu_para.add_run(institution)
            if dates:
                edu_para.add_run(f"  ({dates})")
            if gpa:
                doc.add_paragraph(f"GPA: {gpa}")

    # Skills
    skills = resume_data.get("skills", [])
    if skills:
        _add_section_header(doc, "SKILLS")
        doc.add_paragraph(", ".join(skills))

    # Certifications
    certifications = resume_data.get("certifications", [])
    if certifications:
        _add_section_header(doc, "CERTIFICATIONS")
        for cert in certifications:
            doc.add_paragraph(cert, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_section_header(doc, title: str):
    """Add a formatted section header to the DOCX document."""
    from docx.shared import Pt, RGBColor

    para = doc.add_paragraph()
    run = para.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    # Add a thin border below (via paragraph border)
    # Simple approach: add a line of underscores isn't ATS-friendly
    # Instead we just rely on the bold heading style


def _escape_xml(text: str) -> str:
    """Escape XML special characters for reportlab Paragraph tags."""
    text = text.replace("&", "&amp;")
    text = text.replace("<", "&lt;")
    text = text.replace(">", "&gt;")
    return text
